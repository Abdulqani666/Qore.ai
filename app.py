"""Qore server: serves the page and runs AI jobs in the background.

Keys never reach the browser. Set them as environment variables in Render:
ANTHROPIC_API_KEY, OPENAI_API_KEY, GEMINI_API_KEY and (recommended) APP_PASSWORD.
"""
import hmac
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
import uuid

import requests
from flask import Flask, Response, jsonify, request, send_from_directory

import ai
import db
from prompts import COMMANDS, TAGS, system_for

HERE = os.path.dirname(os.path.abspath(__file__))
app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = 40 * 1024 * 1024

JOBS = {}
LOCK = threading.Lock()
HITS = {}
MAX_TEXT = 60000


def password():
    return ai.env("APP_PASSWORD")


if db.ENABLED:
    try:
        db.init()
    except Exception as _e:      # noqa: BLE001
        print("database setup failed:", _e)


def user_id():
    return db.read_token(request.headers.get("X-Access", "")) if db.ENABLED else None


def authed():
    if db.ENABLED:
        return user_id() is not None
    pw = password()
    return (not pw) or hmac.compare_digest(request.headers.get("X-Access", ""), pw)


def limited():
    ip = (request.headers.get("X-Forwarded-For", request.remote_addr or "") or "").split(",")[0].strip()
    now = time.time()
    h = [t for t in HITS.get(ip, []) if now - t < 3600]
    if len(h) >= int(ai.env("HOURLY_LIMIT", "60")):
        HITS[ip] = h
        return True
    h.append(now)
    HITS[ip] = h
    return False


def guard():
    if not authed():
        return jsonify(error="Wrong or missing password.", auth=True), 401
    return None


# ------------------------------------------------------------ files
def read_file(f):
    """Returns (text, ai_file or None). ai_file is sent to the model as is (images)."""
    name = (f.filename or "file").lower()
    data = f.read()
    mime = f.mimetype or ""
    if name.endswith(".docx"):
        try:
            import docx
            d = docx.Document(io.BytesIO(data))
            out = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for r in t.rows:
                    out.append(" | ".join(c.text.strip() for c in r.cells))
        except Exception:
            # damaged file (for example a bad picture inside): read only the words
            import re, zipfile, html as _h
            try:
                z = zipfile.ZipFile(io.BytesIO(data))
                xml = z.read("word/document.xml").decode("utf8", "ignore")
            except Exception:
                raise ValueError(f"I cannot open {f.filename}. The Word file looks damaged. Open it in Word, save it again, and attach it again.")
            out = []
            for para in re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.S):
                t = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", para, flags=re.S))
                if t.strip():
                    out.append(_h.unescape(t))
        return f"[File {f.filename}]\n" + "\n".join(out), None
    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            txt = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
        except Exception:
            txt = ""
        return f"[File {f.filename}]\n{txt}", None
    if mime.startswith("image/") or name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
        m = mime if mime.startswith("image/") else "image/jpeg"
        return f"[Picture {f.filename} attached]", {"kind": "image", "mime": m, "data": data}
    return f"[File {f.filename}]\n" + data.decode("utf-8", "ignore"), None


# ------------------------------------------------------------ jobs
def new_job():
    jid = uuid.uuid4().hex[:12]
    with LOCK:
        for k in [k for k, v in JOBS.items() if time.time() - v["t"] > 7200]:
            JOBS.pop(k, None)
        JOBS[jid] = {"t": time.time(), "state": "run", "steps": [], "result": None, "error": None}
    return jid


def step(jid, msg):
    with LOCK:
        JOBS[jid]["steps"].append(msg)


def work(jid, cmd, text, tags, history, prev_doc, ai_files):
    spec = COMMANDS[cmd]
    try:
        forced = next((t for t in tags if t in ("claude", "gpt", "gemini")), None)
        tier = "deep" if "deep" in tags else ("fast" if "fast" in tags else None)
        sources = []
        prompt = text
        want_search = spec["search"] and "nosearch" not in tags or "search" in tags
        if want_search and cmd in ("build_document", "lesson", "ask") and text.strip():
            notes, sources = ai.research(text[:1500], say=lambda m: step(jid, m))
            if notes:
                prompt = f"{text}\n\n--- RESEARCH NOTES (from the web; use as sources) ---\n{notes}"
        if cmd == "revise":
            prompt = f"CURRENT DOCUMENT:\n{prev_doc}\n\nCHANGE REQUEST:\n{text}"
        messages = [m for m in history if m.get("text")][-6:] if cmd == "ask" else []
        messages.append({"role": "user", "text": prompt[:MAX_TEXT]})
        step(jid, "Starting...")
        route, think, mx = spec["route"], spec["thinking"] or tier == "deep", 16000 if spec["kind"] == "doc" else 6000
        if cmd == "ask" and not ai_files and not tier and len(text) < 300 and "search" not in tags:
            tier, route, think, mx = "fast", [("gemini", "fast"), ("claude", "fast"), ("gpt", "fast")], False, 1500   # quick chat
        r = ai.run_with_fallback(
            route, system_for(cmd), messages, files=ai_files,
            thinking=think, max_tokens=mx,
            forced=forced, tier_override=tier, say=lambda m: step(jid, m))
        out = {"kind": spec["kind"], "cmd": cmd, "text": r["text"], "model": r["model"],
               "provider": ai.LABEL[r["provider"]], "note": r["note"],
               "sources": (sources + r["sources"])[:8]}
        if spec["kind"] == "json":
            out["data"] = ai.extract_json(r["text"])
        with LOCK:
            JOBS[jid].update(state="done", result=out)
    except ai.AIError as e:
        with LOCK:
            JOBS[jid].update(state="error", error=str(e))
    except Exception as e:      # noqa: BLE001
        with LOCK:
            JOBS[jid].update(state="error", error="Something went wrong: " + str(e)[:160])


# ------------------------------------------------------------ routes
@app.get("/")
def index():
    return send_from_directory(HERE, "index.html")


@app.get("/manifest.webmanifest")
def manifest():
    return send_from_directory(HERE, "manifest.webmanifest", mimetype="application/manifest+json")


@app.get("/sw.js")
def sw():
    r = send_from_directory(HERE, "sw.js", mimetype="text/javascript")
    r.headers["Cache-Control"] = "no-cache"
    return r


@app.get("/icon-<int:n>.png")
def icon(n):
    return send_from_directory(HERE, f"icon-{n}.png") if n in (192, 512) else ("", 404)


@app.get("/healthz")
def healthz():
    return "ok"


@app.get("/api/status")
def status():
    return jsonify(locked=bool(password()) or db.ENABLED, authed=authed(), accounts=db.ENABLED, invite=bool(password()), google=ai.env("GOOGLE_CLIENT_ID"),
                   email=db.email_of(user_id()) if db.ENABLED and user_id() else "",
                   commands={k: {"label": v["label"], "desc": v["desc"], "kind": v["kind"]} for k, v in COMMANDS.items()},
                   tags=TAGS, models=ai.describe() if authed() else {}, fake=ai.FAKE, pdf=soffice_ok())


@app.post("/api/run")
def run():
    g = guard()
    if g:
        return g
    cmd = request.form.get("cmd", "ask")
    if cmd not in COMMANDS or COMMANDS[cmd]["kind"] in ("images", "models"):
        return jsonify(error="Unknown command."), 400
    if limited():
        return jsonify(error="Too many requests. Wait a little and try again."), 429
    text = request.form.get("text", "")
    tags = [t for t in request.form.get("tags", "").split(",") if t in TAGS]
    try:
        history = json.loads(request.form.get("history", "[]"))
        history = [{"role": "assistant" if h.get("role") == "assistant" else "user", "text": str(h.get("text", ""))[:8000]} for h in history]
    except Exception:
        history = []
    prev_doc = request.form.get("prev_doc", "")[:MAX_TEXT]
    texts, ai_files = [], []
    for f in request.files.getlist("files")[:8]:
        try:
            t, af = read_file(f)
        except ValueError as e:
            return jsonify(error=str(e)), 400
        texts.append(t)
        if af:
            ai_files.append(af)
    if texts:
        text = (text + "\n\n" + "\n\n".join(texts)).strip()
    if not text.strip() and not ai_files:
        return jsonify(error="Write something or attach a file first."), 400
    jid = new_job()
    threading.Thread(target=work, args=(jid, cmd, text, tags, history, prev_doc, ai_files), daemon=True).start()
    return jsonify(job=jid)


@app.get("/api/job/<jid>")
def job(jid):
    g = guard()
    if g:
        return g
    with LOCK:
        j = JOBS.get(jid)
        if not j:
            return jsonify(state="lost", error="This job is gone (the server restarted). Please send it again."), 404
        return jsonify(state=j["state"], steps=j["steps"][-6:], result=j["result"], error=j["error"])


@app.get("/api/images")
def images():
    g = guard()
    if g:
        return g
    q = request.args.get("q", "").strip()[:100]
    if not q:
        return jsonify(error="Type what to search for."), 400
    if ai.FAKE:
        return jsonify(items=[{"title": "Test picture", "thumb": "", "url": "https://example.org", "by": "Test", "license": "CC0"}])
    items = []
    hdr = {"User-Agent": "QoreSchoolApp/1.0 (teacher tool)"}
    try:
        r = requests.get("https://api.openverse.org/v1/images/", params={"q": q, "page_size": 12, "license_type": "commercial,modification"},
                         headers=hdr, timeout=20)
        for x in r.json().get("results", []):
            items.append({"title": x.get("title") or q, "thumb": x.get("thumbnail") or x.get("url"), "url": x.get("foreign_landing_url") or x.get("url"),
                          "full": x.get("url"), "by": x.get("creator") or "Unknown", "license": (x.get("license") or "").upper() + " " + (x.get("license_version") or "")})
    except Exception:
        pass
    if len(items) < 6:
        try:
            r = requests.get("https://commons.wikimedia.org/w/api.php", headers=hdr, timeout=20, params={
                "action": "query", "generator": "search", "gsrsearch": q + " filetype:bitmap", "gsrnamespace": 6, "gsrlimit": 12,
                "prop": "imageinfo", "iiprop": "url|extmetadata", "iiurlwidth": 400, "format": "json"})
            for p in (r.json().get("query", {}).get("pages", {}) or {}).values():
                ii = (p.get("imageinfo") or [{}])[0]
                md = ii.get("extmetadata", {})
                items.append({"title": p.get("title", "").replace("File:", ""), "thumb": ii.get("thumburl") or ii.get("url"),
                              "url": ii.get("descriptionurl"), "full": ii.get("url"),
                              "by": (md.get("Artist", {}).get("value", "Wikimedia") or "Wikimedia")[:60],
                              "license": md.get("LicenseShortName", {}).get("value", "")})
        except Exception:
            pass
    if not items:
        return jsonify(error="No pictures found. Try another word."), 404
    return jsonify(items=items[:16])


@app.get("/api/imgfetch")
def imgfetch():
    """Fetches one picture from a trusted free-picture site so the page can put it in a lesson."""
    g = guard()
    if g:
        return g
    from urllib.parse import urlparse
    u = request.args.get("u", "")
    p = urlparse(u)
    host = (p.hostname or "").lower()
    if p.scheme != "https" or not (host.endswith(".wikimedia.org") or host == "api.openverse.org" or host.endswith(".openverse.org")):
        return jsonify(error="This picture site is not allowed."), 400
    if ai.FAKE:
        return jsonify(error="No pictures in test mode."), 404
    try:
        r = requests.get(u, headers={"User-Agent": "QoreSchoolApp/1.0 (teacher tool)"}, timeout=20, stream=True)
        data = r.raw.read(4_000_001, decode_content=True)
        ct = r.headers.get("Content-Type", "")
        if r.status_code != 200 or not ct.startswith("image/") or len(data) > 4_000_000:
            return jsonify(error="Picture not available."), 404
        return Response(data, mimetype=ct)
    except Exception:
        return jsonify(error="Picture not available."), 404


# ------------------------------------------------------------ Word -> PDF (LibreOffice)
CONV = threading.Lock()
PDFS = {}


def soffice_ok():
    return bool(shutil.which("soffice") or shutil.which("libreoffice"))


def convert_job(jid, data):
    try:
        step(jid, "Making the PDF...")
        d = tempfile.mkdtemp(prefix="qpdf")
        src = os.path.join(d, "doc.docx")
        with open(src, "wb") as f:
            f.write(data)
        exe = shutil.which("soffice") or shutil.which("libreoffice")
        with CONV:
            subprocess.run([exe, "--headless", "--norestore", "-env:UserInstallation=file:///tmp/lo_profile",
                            "--convert-to", "pdf", "--outdir", d, src], timeout=150, capture_output=True)
        pdf = os.path.join(d, "doc.pdf")
        if not os.path.exists(pdf):
            raise RuntimeError("The PDF could not be made from this file.")
        pages = 1
        try:
            out = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True, timeout=30).stdout
            pages = int(next(l for l in out.splitlines() if l.startswith("Pages:")).split()[1])
        except Exception:
            pass
        pid = uuid.uuid4().hex[:12]
        with LOCK:
            for k in [k for k, v in PDFS.items() if time.time() - v["t"] > 7200]:
                shutil.rmtree(PDFS.pop(k)["dir"], ignore_errors=True)
            PDFS[pid] = {"t": time.time(), "dir": d, "pdf": pdf, "pages": pages}
            JOBS[jid].update(state="done", result={"id": pid, "pages": pages})
    except Exception as e:      # noqa: BLE001
        with LOCK:
            JOBS[jid].update(state="error", error=str(e)[:160])


@app.post("/api/pdf")
def pdf_start():
    g = guard()
    if g:
        return g
    if not soffice_ok():
        return jsonify(error="PDF maker is not installed on this server."), 501
    if limited():
        return jsonify(error="Too many requests. Wait a little and try again."), 429
    f = request.files.get("file")
    if not f:
        return jsonify(error="No file."), 400
    jid = new_job()
    threading.Thread(target=convert_job, args=(jid, f.read()), daemon=True).start()
    return jsonify(job=jid)


@app.post("/api/convert")
def convert_to_docx():
    """Turns a LibreOffice/old Word file (odt, ott, doc, rtf) into .docx so Qore can read the template."""
    g = guard()
    if g:
        return g
    if not soffice_ok():
        return jsonify(error="LibreOffice is not installed on this server."), 501
    if limited():
        return jsonify(error="Too many requests. Wait a little and try again."), 429
    f = request.files.get("file")
    if not f:
        return jsonify(error="No file."), 400
    ext = os.path.splitext(f.filename or "")[1].lower()
    if ext not in (".odt", ".ott", ".doc", ".dot", ".rtf", ".dotx", ".dotm", ".docm"):
        return jsonify(error="Use a Word or LibreOffice file."), 400
    d = tempfile.mkdtemp(prefix="qconv")
    try:
        src = os.path.join(d, "in" + ext)
        f.save(src)
        exe = shutil.which("soffice") or shutil.which("libreoffice")
        with CONV:
            subprocess.run([exe, "--headless", "--norestore", "-env:UserInstallation=file:///tmp/lo_profile",
                            "--convert-to", "docx:MS Word 2007 XML", "--outdir", d, src], timeout=120, capture_output=True)
        out = os.path.join(d, "in.docx")
        if not os.path.exists(out):
            return jsonify(error="I could not open this file."), 422
        with open(out, "rb") as fh:
            data = fh.read()
        return Response(data, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    finally:
        shutil.rmtree(d, ignore_errors=True)


@app.get("/api/pdf/<pid>")
def pdf_get(pid):
    g = guard()
    if g:
        return g
    p = PDFS.get(pid)
    if not p:
        return jsonify(error="This PDF is gone. Make it again."), 404
    with open(p["pdf"], "rb") as f:
        return Response(f.read(), mimetype="application/pdf")


@app.get("/api/pdf/<pid>/page/<int:n>")
def pdf_page(pid, n):
    g = guard()
    if g:
        return g
    p = PDFS.get(pid)
    if not p or n < 1 or n > p["pages"]:
        return jsonify(error="No such page."), 404
    base = os.path.join(p["dir"], f"pg{n}")
    if not os.path.exists(base + ".png"):
        subprocess.run(["pdftoppm", "-r", "110", "-png", "-singlefile", "-f", str(n), "-l", str(n), p["pdf"], base],
                       capture_output=True, timeout=60)
    if not os.path.exists(base + ".png"):
        return jsonify(error="Could not draw the page."), 500
    with open(base + ".png", "rb") as f:
        return Response(f.read(), mimetype="image/png")


# ------------------------------------------------------------ accounts, history, files
AUTH_HITS = {}


def auth_limited():
    ip = (request.headers.get("X-Forwarded-For", request.remote_addr or "") or "").split(",")[0].strip()
    now = time.time()
    h = [t for t in AUTH_HITS.get(ip, []) if now - t < 3600]
    AUTH_HITS[ip] = h + [now]
    return len(h) >= 25


@app.post("/api/auth")
def auth():
    if not db.ENABLED:
        return jsonify(error="Accounts are not set up on this server."), 501
    if auth_limited():
        return jsonify(error="Too many tries. Wait a while and try again."), 429
    j = request.get_json(silent=True) or {}
    try:
        if j.get("mode") == "signup":
            pw = password()
            if pw and not hmac.compare_digest(str(j.get("code", "")), pw):
                return jsonify(error="The access code is wrong. Ask the owner of this app for it."), 403
            uid, err = db.signup(j.get("email"), j.get("password"))
        else:
            uid, err = db.login(j.get("email"), j.get("password"))
    except Exception as e:      # noqa: BLE001
        return jsonify(error="The database is not ready yet. Try again in a minute. " + str(e)[:80]), 503
    if err:
        return jsonify(error=err), 400
    return jsonify(token=db.make_token(uid), email=db.email_of(uid))


@app.get("/api/diagnose")
def diagnose():
    if not authed():
        return jsonify(error="Please sign in.", auth=True), 401
    import concurrent.futures as cf
    on = ai.available()

    def one(p):
        t0 = time.time()
        try:
            ids = ai._listed(p)
            _t, model, _s = ai.call(p, "fast", "Reply with the single word OK.", [{"role": "user", "text": "Say OK"}], max_tokens=30)
            return p, {"ok": True, "model": model, "seconds": round(time.time() - t0, 1)}
        except Exception as e:      # noqa: BLE001
            return p, {"ok": False, "error": str(e)[:300], "seconds": round(time.time() - t0, 1),
                       "seen": (len(ids) if ids else 0) if "ids" in dir() else 0}
    out = {}
    with cf.ThreadPoolExecutor(3) as ex:
        for p, r in ex.map(one, [p for p in ("claude", "gpt", "gemini") if on.get(p)]):
            out[p] = r
    return jsonify(results=out, labels=ai.LABEL)


@app.post("/api/auth/google")
def auth_google():
    cid = ai.env("GOOGLE_CLIENT_ID")
    if not db.ENABLED or not cid:
        return jsonify(error="Google sign-in is not set up on this server."), 501
    if auth_limited():
        return jsonify(error="Too many tries. Wait a while and try again."), 429
    j = request.get_json(silent=True) or {}
    cred = str(j.get("credential", ""))
    try:
        if ai.FAKE and cred.startswith("fake:"):
            email = cred[5:]
        else:
            r = requests.get("https://oauth2.googleapis.com/tokeninfo", params={"id_token": cred}, timeout=15)
            info = r.json() if r.ok else {}
            if info.get("aud") != cid or str(info.get("email_verified")).lower() != "true" or not info.get("email"):
                return jsonify(error="Google could not confirm this account. Try again."), 400
            email = info["email"]
        uid, new = db.google_user(email)
        if new:
            pw = password()
            if pw and not hmac.compare_digest(str(j.get("code", "")), pw):
                db.run("DELETE FROM users WHERE id=%s", (uid,))
                return jsonify(error="New accounts need the access code from the school. Type it in the box, then tap Google again.", need_code=True), 403
    except Exception as e:      # noqa: BLE001
        return jsonify(error="Google sign-in failed. " + str(e)[:80]), 503
    return jsonify(token=db.make_token(uid), email=db.email_of(uid))


def need_user():
    if not db.ENABLED:
        return None, (jsonify(error="History needs sign-in, which is not set up on this server."), 501)
    uid = user_id()
    if uid is None:
        return None, (jsonify(error="Please sign in.", auth=True), 401)
    return uid, None


@app.get("/api/chats")
def chats_list():
    uid, err = need_user()
    if err:
        return err
    rows = db.run("SELECT id,title,EXTRACT(EPOCH FROM updated) FROM chats WHERE uid=%s ORDER BY updated DESC LIMIT 200", (uid,), "all")
    return jsonify(chats=[{"id": r[0], "title": r[1], "updated": int(r[2])} for r in rows])


@app.get("/api/chats/<cid>")
def chats_get(cid):
    uid, err = need_user()
    if err:
        return err
    r = db.run("SELECT title,items FROM chats WHERE id=%s AND uid=%s", (cid, uid), "one")
    if not r:
        return jsonify(error="Not found."), 404
    return jsonify(id=cid, title=r[0], items=json.loads(r[1]))


@app.put("/api/chats/<cid>")
def chats_put(cid):
    uid, err = need_user()
    if err:
        return err
    if not re.fullmatch(r"[0-9a-f]{8,32}", cid):
        return jsonify(error="Bad id."), 400
    j = request.get_json(silent=True) or {}
    items = json.dumps(j.get("items", []))[:3_000_000]
    if not items.endswith("]"):
        return jsonify(error="Chat is too long to save."), 413
    owner = db.run("SELECT uid FROM chats WHERE id=%s", (cid,), "one")
    if owner and owner[0] != uid:
        return jsonify(error="Not yours."), 403
    db.run("""INSERT INTO chats(id,uid,title,items,updated) VALUES(%s,%s,%s,%s,now())
              ON CONFLICT (id) DO UPDATE SET title=EXCLUDED.title, items=EXCLUDED.items, updated=now()""",
           (cid, uid, str(j.get("title", ""))[:120], items))
    return jsonify(ok=True)


@app.delete("/api/chats/<cid>")
def chats_del(cid):
    uid, err = need_user()
    if err:
        return err
    db.run("DELETE FROM chats WHERE id=%s AND uid=%s", (cid, uid))
    return jsonify(ok=True)


@app.put("/api/files/<fid>")
def files_put(fid):
    uid, err = need_user()
    if err:
        return err
    if not re.fullmatch(r"[0-9a-f]{8,32}", fid):
        return jsonify(error="Bad id."), 400
    data = request.get_data()
    if not data or len(data) > 12_000_000:
        return jsonify(error="File is empty or too big."), 400
    owner = db.run("SELECT uid FROM files WHERE id=%s", (fid,), "one")
    if owner and owner[0] != uid:
        return jsonify(error="Not yours."), 403
    import psycopg2
    db.run("""INSERT INTO files(id,uid,name,data) VALUES(%s,%s,%s,%s)
              ON CONFLICT (id) DO UPDATE SET data=EXCLUDED.data, name=EXCLUDED.name""",
           (fid, uid, request.args.get("name", "file.docx")[:200], psycopg2.Binary(data)))
    return jsonify(ok=True)


@app.get("/api/files/<fid>")
def files_get(fid):
    uid, err = need_user()
    if err:
        return err
    r = db.run("SELECT data FROM files WHERE id=%s AND uid=%s", (fid, uid), "one")
    if not r:
        return jsonify(error="This file is gone."), 404
    return Response(bytes(r[0]), mimetype="application/octet-stream")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")))
